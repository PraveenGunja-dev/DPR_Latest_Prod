from pptx import Presentation

def get_slide_titles(pptx_path):
    prs = Presentation(pptx_path)
    for i, slide in enumerate(prs.slides):
        title = "No Title"
        if slide.shapes.title and slide.shapes.title.text:
            title = slide.shapes.title.text
        elif slide.shapes:
            # Fallback to the first text shape if there is no explicit title shape
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    title = shape.text.strip().replace("\n", " ").replace("\r", "")[:50] + "..."
                    break
        print(f"Slide {i+1}: {title}")

if __name__ == "__main__":
    get_slide_titles(r"d:\BESS\data\CPAG Deck.pptx")
